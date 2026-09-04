import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SAMPLE_PATH = PROJECT_ROOT / "samples" / "fictional_hmi_test_spec.xlsx"
EXPECTED_REGIONS_PATH = (
    PROJECT_ROOT
    / "tests"
    / "fixtures"
    / "complex_documents"
    / "fictional_hmi_expected_regions.json"
)


def test_fictional_hmi_fixture_matches_the_day1_design() -> None:
    """验证架空 HMI 样本的 Sheet、结构范围和关键文本。 / Verifies the fictional HMI sample's sheets, regions, and key text."""

    expected = json.loads(EXPECTED_REGIONS_PATH.read_text(encoding="utf-8"))
    assert expected["fictional_data"] is True
    assert expected["source_file"] == SAMPLE_PATH.name

    workbook = load_workbook(SAMPLE_PATH, data_only=True)
    try:
        assert workbook.sheetnames == ["README", "機能仕様", "画面遷移", "CAN信号"]

        feature_sheet = workbook["機能仕様"]
        assert feature_sheet["E8"].value == "CAN信号"
        assert feature_sheet["E9"].value == "信号名"
        assert feature_sheet["A10"].value == "HMI-AC-001"
        assert feature_sheet["F10"].value == "01h"
        assert "A10:A11" in {str(item) for item in feature_sheet.merged_cells.ranges}

        for region in expected["expected_regions"]:
            worksheet = workbook[region["sheet"]]
            min_column, min_row, max_column, max_row = range_boundaries(
                region["cell_range"]
            )
            values = [
                worksheet.cell(row=row, column=column).value
                for row in range(min_row, max_row + 1)
                for column in range(min_column, max_column + 1)
            ]
            assert any(value not in (None, "") for value in values), region["region_id"]
    finally:
        workbook.close()
