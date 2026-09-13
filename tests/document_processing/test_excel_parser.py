from openpyxl import Workbook
from openpyxl.styles import Border, Side

from app.document_processing.parsers import ExcelDocumentParser, PdfDocumentParser


def test_excel_parser_preserves_sheet_headers_rows_and_merged_values(tmp_path) -> None:
    """验证 Excel 解析保留 Sheet、Header、Row 和合并单元格语义。 / Verifies Excel parsing preserves sheet, header, row, and merged-cell semantics."""

    path = tmp_path / "travel_policy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "国内出張"
    sheet.append(["区分", "上限", "備考"])
    sheet.append(["宿泊費", 10000, "1泊"])
    sheet.merge_cells("A3:A4")
    sheet["A3"] = "交通費"
    sheet["B3"] = "実費"
    sheet["B4"] = "領収書必須"

    thin_side = Side(style="thin")
    table_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    for row in sheet.iter_rows(min_row=1, max_row=4, min_col=1, max_col=3):
        for cell in row:
            cell.border = table_border

    workbook.save(path)
    workbook.close()

    blocks = ExcelDocumentParser().parse(path)

    assert len(blocks) == 1
    assert blocks[0].content_type == "table"
    assert blocks[0].sheet == "国内出張"
    assert blocks[0].cell_range == "A1:C4"
    assert blocks[0].rows == "2:4"
    assert "区分=宿泊費" in blocks[0].content
    assert "上限=10000" in blocks[0].content
    assert blocks[0].content.count("区分=交通費") == 2


def test_excel_parser_keeps_unbordered_multiple_lines_as_paragraph(tmp_path) -> None:
    """验证无边框连续多行不会因为行数大于一而误判为表格。 / Verifies unbordered lines are not misclassified as a table."""

    path = tmp_path / "unbordered_notes.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "説明"
    sheet["A1"] = "利用条件"
    sheet["A3"] = "申請は三日前までに行います。"
    sheet["A4"] = "上司の承認が必要です。"
    sheet["A5"] = "承認後に休暇を取得できます。"
    workbook.save(path)
    workbook.close()

    blocks = ExcelDocumentParser().parse(path)

    paragraph_blocks = []
    for block in blocks:
        if block.content_type == "paragraph":
            paragraph_blocks.append(block)

    assert len(paragraph_blocks) == 1
    assert paragraph_blocks[0].cell_range == "A3:A5"
    assert "申請は三日前までに行います。" in paragraph_blocks[0].content
    assert "承認後に休暇を取得できます。" in paragraph_blocks[0].content


def test_excel_parser_requires_a_closed_border_for_table(tmp_path) -> None:
    """验证超过两个单元格且闭合的边框矩形才被识别为表格。 / Verifies a closed border rectangle with more than two cells is a table."""

    path = tmp_path / "bordered_table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "残日数"
    sheet.append(["社員", "残日数"])
    sheet.append(["U001", 8])

    thin_side = Side(style="thin")
    table_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    for row in sheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=2):
        for cell in row:
            cell.border = table_border

    workbook.save(path)
    workbook.close()

    blocks = ExcelDocumentParser().parse(path)

    assert len(blocks) == 1
    assert blocks[0].content_type == "table"
    assert blocks[0].cell_range == "A1:B2"


def test_excel_parser_rejects_an_open_border_as_table(tmp_path) -> None:
    """验证边框没有闭合时不能把区域判断为表格。 / Verifies an open border is not classified as a table."""

    path = tmp_path / "open_border.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "説明"
    sheet["A1"] = "項目"
    sheet["B1"] = "内容"
    sheet["A2"] = "申請"
    sheet["B2"] = "三日前"

    thin_side = Side(style="thin")
    almost_closed_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
    )
    for row in sheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=2):
        for cell in row:
            cell.border = almost_closed_border

    workbook.save(path)
    workbook.close()

    blocks = ExcelDocumentParser().parse(path)

    assert len(blocks) == 1
    assert blocks[0].content_type == "paragraph"


def test_excel_parser_adds_nearby_title_and_note_to_table_content(tmp_path) -> None:
    """验证表格附近的标题和备注会一起进入表格检索正文。 / Verifies nearby titles and notes enrich table search text."""

    path = tmp_path / "table_context.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "年休"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "表1：年休残日数"
    sheet.append(["社員", "年度", "残日数"])
    sheet.append(["U001", 2026, 8])
    sheet.append(["U002", 2026, 5])
    sheet["A5"] = "備考：残日数は申請承認後に更新されます。"

    thin_side = Side(style="thin")
    table_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    for row in sheet.iter_rows(min_row=2, max_row=4, min_col=1, max_col=3):
        for cell in row:
            cell.border = table_border

    workbook.save(path)
    workbook.close()

    blocks = ExcelDocumentParser().parse(path)

    table_block = None
    for block in blocks:
        if block.content_type == "table":
            table_block = block
            break

    assert table_block is not None
    assert "表1：年休残日数" in table_block.content
    assert "備考：残日数は申請承認後に更新されます。" in table_block.content
    assert table_block.cell_range == "A2:C4"


def test_excel_parser_separates_an_unbordered_description_before_a_table(
    tmp_path,
) -> None:
    """验证紧贴表格的无边框说明会独立解析并加入表格上下文。 / Verifies an adjacent unbordered description is separated and added to table context."""

    path = tmp_path / "description_before_table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "年休"
    sheet["A1"] = "年休データ"
    sheet["A3"] = "以下の表は2026年度の残日数を示します。"
    sheet["A4"] = "社員"
    sheet["B4"] = "残日数"
    sheet["A5"] = "U001"
    sheet["B5"] = 8

    thin_side = Side(style="thin")
    table_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    for row in sheet.iter_rows(min_row=4, max_row=5, min_col=1, max_col=2):
        for cell in row:
            cell.border = table_border

    workbook.save(path)
    workbook.close()

    blocks = ExcelDocumentParser().parse(path)

    table_block = None
    for block in blocks:
        if block.content_type == "table":
            table_block = block
            break

    assert table_block is not None
    assert "以下の表は2026年度の残日数を示します。" in table_block.content
    assert table_block.cell_range == "A4:B5"


class FakePdfPage:
    """为 PDF Parser 测试提供固定页面文本。 / Provides fixed page text for PDF parser tests."""

    def __init__(self, text: str) -> None:
        """保存页面提取结果。 / Stores the page extraction result."""

        self._text = text

    def extract_text(self) -> str:
        """返回固定文本层内容。 / Returns fixed text-layer content."""

        return self._text


class FakePdfReader:
    """模拟包含两个文本页的 pypdf Reader。 / Simulates a pypdf reader containing two text pages."""

    def __init__(self, _path) -> None:
        """创建固定测试页面。 / Creates fixed test pages."""

        self.pages = [
            FakePdfPage("第一段。\n\n第二段。"),
            FakePdfPage("第三页内容。"),
        ]


def test_pdf_parser_preserves_page_numbers_and_paragraph_chunks(
    monkeypatch,
    tmp_path,
) -> None:
    """验证 PDF Parser 保留页码并按自然段组合 Chunk。 / Verifies the PDF parser preserves page numbers and groups natural paragraphs."""

    monkeypatch.setattr(
        "app.document_processing.pdf_parser.PdfReader",
        FakePdfReader,
    )

    blocks = PdfDocumentParser(maximum_chunk_characters=100).parse(
        tmp_path / "policy.pdf"
    )

    assert [block.page for block in blocks] == [1, 2]
    assert blocks[0].content_type == "paragraph"
    assert blocks[0].content == "第一段。\n第二段。"
    assert blocks[1].content == "第三页内容。"
