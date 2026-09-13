from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PillowImage

from app.document_processing.excel_image_extractor import ExcelImageExtractor
from app.document_processing.excel_parser import ExcelDocumentParser
from app.document_processing.storage import LocalImageAssetStorage


def create_excel_with_images(workbook_path: Path, image_path: Path) -> None:
    """创建包含两个 Sheet 和两张嵌入图片的虚构 Excel。 / Creates a fictional workbook with two sheets and two embedded images."""

    image = PillowImage.new("RGB", (8, 8), color=(25, 90, 180))
    image.save(image_path, format="PNG")
    image.close()

    workbook = Workbook()
    first_sheet = workbook.active
    first_sheet.title = "System"
    first_sheet["A1"] = "系统结构图"
    first_sheet.add_image(OpenpyxlImage(str(image_path)), "B2")

    second_sheet = workbook.create_sheet("Operation")
    second_sheet["A1"] = "操作画面"
    second_sheet.add_image(OpenpyxlImage(str(image_path)), "D5")

    workbook.save(workbook_path)
    workbook.close()


def test_excel_image_extractor_saves_images_with_source_metadata(
    tmp_path: Path,
) -> None:
    """验证嵌入图片被保存，并保留 Sheet、序号、MIME 和可靠锚点。 / Verifies embedded images are stored with sheet, index, MIME, and reliable anchors."""

    workbook_path = tmp_path / "hmi_spec.xlsx"
    source_image_path = tmp_path / "diagram.png"
    create_excel_with_images(workbook_path, source_image_path)
    storage = LocalImageAssetStorage(tmp_path / "document_storage")
    extractor = ExcelImageExtractor(storage)

    images = extractor.extract(
        workbook_path,
        document_id="HMI-SPEC",
        document_version_id="HMI-SPEC-V1",
    )

    assert len(images) == 2

    first_image = images[0]
    assert first_image.sheet == "System"
    assert first_image.cell_range == "B2"
    assert first_image.image_index == 1
    assert first_image.mime_type == "image/png"
    assert first_image.image_path.is_file()

    second_image = images[1]
    assert second_image.sheet == "Operation"
    assert second_image.cell_range == "D5"
    assert second_image.image_index == 2
    assert second_image.image_path.is_file()
    assert second_image.image_id != first_image.image_id


def test_excel_image_id_is_stable_when_extraction_is_repeated(tmp_path: Path) -> None:
    """验证重复提取同一 Workbook 时 image_id 和位置保持稳定。 / Verifies repeated extraction keeps image IDs and locations stable."""

    workbook_path = tmp_path / "hmi_spec.xlsx"
    source_image_path = tmp_path / "diagram.png"
    create_excel_with_images(workbook_path, source_image_path)
    extractor = ExcelImageExtractor(LocalImageAssetStorage(tmp_path / "storage"))

    first_result = extractor.extract(workbook_path, "HMI-SPEC", "HMI-SPEC-V1")
    second_result = extractor.extract(workbook_path, "HMI-SPEC", "HMI-SPEC-V1")

    assert first_result[0].image_id == second_result[0].image_id
    assert first_result[0].cell_range == second_result[0].cell_range
    assert first_result[1].image_id == second_result[1].image_id
    assert first_result[1].cell_range == second_result[1].cell_range


def test_excel_without_images_returns_empty_result(tmp_path: Path) -> None:
    """验证没有嵌入图片的 Excel 返回空列表。 / Verifies a workbook without embedded images returns an empty list."""

    workbook_path = tmp_path / "text_only.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "纯文字"
    workbook.save(workbook_path)
    workbook.close()

    extractor = ExcelImageExtractor(LocalImageAssetStorage(tmp_path / "storage"))

    images = extractor.extract(workbook_path, "TEXT-DOC", "TEXT-DOC-V1")

    assert images == []


def test_existing_excel_text_parser_still_reads_workbook_with_images(
    tmp_path: Path,
) -> None:
    """验证新增提取器不会改变现有 Excel 文字 ParsedBlock。 / Verifies the new extractor does not change existing Excel text ParsedBlocks."""

    workbook_path = tmp_path / "hmi_spec.xlsx"
    source_image_path = tmp_path / "diagram.png"
    create_excel_with_images(workbook_path, source_image_path)

    blocks = ExcelDocumentParser().parse(workbook_path)

    assert len(blocks) == 2
    assert blocks[0].content == "系统结构图"
    assert blocks[0].sheet == "System"
    assert blocks[1].content == "操作画面"
    assert blocks[1].sheet == "Operation"
