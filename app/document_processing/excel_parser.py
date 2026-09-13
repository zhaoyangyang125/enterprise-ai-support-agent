"""把 Excel 文档转换成可以检索的内容块。

初学者建议按照下面的顺序阅读：

1. ``ExcelDocumentParser.parse``：打开 Excel，并逐个处理 Sheet。
2. ``_parse_sheet``：从上往下读取一张 Sheet。
3. ``_find_standalone_block``：识别单独占一行的标题、章节和备注。
4. ``_build_region_block``：把连续的普通行识别成键值对、段落或表格。
5. 其他 ``_build_*`` 方法：负责生成最终可以检索的文字。

This module converts an Excel document into searchable content blocks.
The numbered list above is the recommended reading order for beginners.
"""

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.document import ParsedBlock


_EMPTY_VALUES = (None, "")
_NOTE_PREFIXES = ("注意", "備考", "解析上の期待", "注記", "補足")
_SECTION_PATTERN = re.compile(r"^\s*(?:\d+(?:\.\d+)*[.．]?\s+|第.+[章節])")


class _WorksheetReader:
    """统一读取普通单元格和合并单元格。 / Reads normal and merged cells consistently.

    openpyxl 只把合并区域的值保存在左上角单元格中。例如 A3:A4
    合并以后，A3 有值而 A4 是空值。本类让 A3 和 A4 都能读取到相同的
    “逻辑值”，但不会修改原来的 Workbook。
    """

    def __init__(self, worksheet: Worksheet) -> None:
        """保存工作表，并准备合并单元格的逻辑值。 / Stores a sheet and prepares logical merged-cell values."""

        self.worksheet = worksheet
        self.merged_ranges = [
            CellRange(str(merged_range))
            for merged_range in worksheet.merged_cells.ranges
        ]
        self._merged_values = self._build_merged_value_map()

    def _build_merged_value_map(self) -> dict[tuple[int, int], Any]:
        """记录每个合并单元格应该读取到的值。 / Maps every merged cell to its logical value."""

        merged_values: dict[tuple[int, int], Any] = {}

        for merged_range in self.merged_ranges:
            top_left_value = self.worksheet.cell(
                row=merged_range.min_row,
                column=merged_range.min_col,
            ).value

            for row_number in range(
                merged_range.min_row,
                merged_range.max_row + 1,
            ):
                for column_number in range(
                    merged_range.min_col,
                    merged_range.max_col + 1,
                ):
                    merged_values[(row_number, column_number)] = top_left_value

        return merged_values

    def get_value(self, row_number: int, column_number: int) -> Any:
        """读取单元格的逻辑值。 / Returns the logical value of a cell."""

        position = (row_number, column_number)
        if position in self._merged_values:
            return self._merged_values[position]

        return self.worksheet.cell(
            row=row_number,
            column=column_number,
        ).value

    def get_real_nonempty_cells(self, row_number: int) -> list[tuple[int, Any]]:
        """读取一行中真正保存了值的单元格。 / Returns physically populated cells in a row.

        这里故意不展开合并值。这个方法主要用于判断某一行是不是只有一个
        标题，或者是不是由 Label/Value 组成的键值对。
        """

        cells: list[tuple[int, Any]] = []

        for column_number in range(1, self.worksheet.max_column + 1):
            value = self.worksheet.cell(
                row=row_number,
                column=column_number,
            ).value
            if value not in _EMPTY_VALUES:
                cells.append((column_number, value))

        return cells

    def find_merged_range_starting_at(
        self,
        row_number: int,
        column_number: int,
    ) -> CellRange | None:
        """查找从指定单元格开始的合并区域。 / Finds a merged range starting at a cell."""

        for merged_range in self.merged_ranges:
            starts_at_requested_cell = (
                merged_range.min_row == row_number
                and merged_range.min_col == column_number
            )
            if starts_at_requested_cell:
                return merged_range

        return None

    def is_empty_row(self, row_number: int) -> bool:
        """判断一整行是否为空。 / Checks whether an entire row is empty."""

        for column_number in range(1, self.worksheet.max_column + 1):
            if self.get_value(row_number, column_number) not in _EMPTY_VALUES:
                return False

        return True

    def find_used_column_bounds(self, rows: list[int]) -> tuple[int, int]:
        """返回一组行实际使用的最小列和最大列。 / Returns the used column bounds for rows."""

        used_columns: list[int] = []

        for row_number in rows:
            for column_number in range(1, self.worksheet.max_column + 1):
                value = self.get_value(row_number, column_number)
                if value not in _EMPTY_VALUES:
                    used_columns.append(column_number)

        return min(used_columns), max(used_columns)


class ExcelDocumentParser:
    """将 Excel 的内容区域转换为可定位的 ParsedBlock。 / Converts Excel regions into located ParsedBlocks."""

    def parse(self, path: Path) -> list[ParsedBlock]:
        """打开 Excel，按顺序解析其中的每一个 Sheet。 / Opens an Excel file and parses every sheet."""

        workbook = load_workbook(path, data_only=True)

        try:
            all_blocks: list[ParsedBlock] = []

            for worksheet in workbook.worksheets:
                sheet_blocks = self._parse_sheet(worksheet)
                all_blocks.extend(sheet_blocks)

            return all_blocks
        finally:
            # 即使解析途中发生异常，也要关闭文件，避免文件一直被占用。
            # Always closes the file, even when parsing fails.
            workbook.close()

    def _parse_sheet(self, worksheet: Worksheet) -> list[ParsedBlock]:
        """从上往下读取一张 Sheet，并按区域生成 Block。 / Reads one sheet from top to bottom and builds blocks."""

        reader = _WorksheetReader(worksheet)
        blocks: list[ParsedBlock] = []

        # 普通的连续非空行先暂存在这里，遇到空行或独立标题后再一起转换。
        # Ordinary consecutive rows are buffered until a separator is found.
        pending_rows: list[int] = []
        current_section: str | None = None
        row_number = 1

        while row_number <= worksheet.max_row:
            standalone_result = self._find_standalone_block(
                reader,
                row_number,
                current_section,
            )

            # 情况一：当前行是单独的标题、章节或备注。
            if standalone_result is not None:
                pending_blocks = self._build_region_blocks(
                    reader,
                    pending_rows,
                    current_section,
                )
                blocks.extend(pending_blocks)
                pending_rows = []

                standalone_block, last_row, starts_new_section = standalone_result
                blocks.append(standalone_block)

                if starts_new_section:
                    current_section = standalone_block.content

                # 合并标题可能占据多行，所以直接跳到合并区域的下一行。
                row_number = last_row + 1
                continue

            # 情况二：空行表示前面的普通内容区域结束。
            if reader.is_empty_row(row_number):
                pending_blocks = self._build_region_blocks(
                    reader,
                    pending_rows,
                    current_section,
                )
                blocks.extend(pending_blocks)
                pending_rows = []

            # 情况三：普通非空行，暂存起来等待组成一个区域。
            else:
                pending_rows.append(row_number)

            row_number += 1

        # 文件末尾不一定有空行，因此循环结束后还要处理最后一个区域。
        final_blocks = self._build_region_blocks(
            reader,
            pending_rows,
            current_section,
        )
        blocks.extend(final_blocks)

        return self._attach_table_context(blocks)

    def _build_region_blocks(
        self,
        reader: _WorksheetReader,
        rows: list[int],
        current_section: str | None,
    ) -> list[ParsedBlock]:
        """按边框变化拆分连续行，再分别建立内容块。 / Splits consecutive rows when border usage changes and builds blocks."""

        if not rows:
            return []

        row_groups: list[list[int]] = []
        current_group: list[int] = []
        current_group_has_border: bool | None = None

        for row_number in rows:
            row_has_border = self._row_has_visible_border(reader, row_number)

            starts_new_group = (
                current_group_has_border is not None
                and row_has_border != current_group_has_border
            )
            if starts_new_group:
                row_groups.append(current_group)
                current_group = []

            current_group.append(row_number)
            current_group_has_border = row_has_border

        if current_group:
            row_groups.append(current_group)

        blocks: list[ParsedBlock] = []
        for row_group in row_groups:
            block = self._build_region_block(
                reader,
                row_group,
                current_section,
            )
            if block is not None:
                blocks.append(block)

        return blocks

    def _row_has_visible_border(
        self,
        reader: _WorksheetReader,
        row_number: int,
    ) -> bool:
        """判断一行是否包含至少一条可见单元格边框。 / Checks whether a row contains any visible cell border."""

        for column_number in range(1, reader.worksheet.max_column + 1):
            cell = reader.worksheet.cell(
                row=row_number,
                column=column_number,
            )
            for side_name in ("left", "right", "top", "bottom"):
                if self._has_visible_border(cell, side_name):
                    return True

        return False

    def _find_standalone_block(
        self,
        reader: _WorksheetReader,
        row_number: int,
        current_section: str | None,
    ) -> tuple[ParsedBlock, int, bool] | None:
        """识别单独占一行的标题、章节或备注。 / Detects a standalone title, section, or note."""

        cells = reader.get_real_nonempty_cells(row_number)

        # 单独标题只能有一个真正保存值的单元格。
        if len(cells) != 1:
            return None

        column_number, raw_value = cells[0]
        text = self._to_text(raw_value)
        merged_range = reader.find_merged_range_starting_at(
            row_number,
            column_number,
        )

        occupies_multiple_cells = False
        if merged_range is not None:
            occupies_multiple_cells = (
                merged_range.max_col > merged_range.min_col
                or merged_range.max_row > merged_range.min_row
            )

        is_note = text.startswith(_NOTE_PREFIXES)
        is_section = bool(_SECTION_PATTERN.match(text))
        is_title = occupies_multiple_cells or row_number <= 2 or is_section

        if not is_note and not is_title:
            return None

        if merged_range is None:
            merged_range = CellRange(
                min_col=column_number,
                min_row=row_number,
                max_col=column_number,
                max_row=row_number,
            )

        if is_note:
            content_type = "note"
            section = current_section
        else:
            content_type = "title"
            section = None

        block = ParsedBlock(
            content=text,
            content_type=content_type,
            section=section,
            sheet=reader.worksheet.title,
            cell_range=str(merged_range),
            rows=self._format_row_range(
                merged_range.min_row,
                merged_range.max_row,
            ),
        )

        return block, merged_range.max_row, is_section

    def _build_region_block(
        self,
        reader: _WorksheetReader,
        rows: list[int],
        current_section: str | None,
    ) -> ParsedBlock | None:
        """把连续的普通行转换为键值对、段落或表格。 / Converts consecutive rows into a key-value block, paragraph, or table."""

        if not rows:
            return None

        first_column, last_column = reader.find_used_column_bounds(rows)
        cell_range = self._format_cell_range(
            first_column,
            rows[0],
            last_column,
            rows[-1],
        )

        # 第一种：Label/Value、Label/Value 形式的键值对区域。
        if self._is_key_value_region(reader, rows):
            return ParsedBlock(
                content=self._build_key_value_text(reader, rows),
                content_type="key_value",
                section=current_section,
                sheet=reader.worksheet.title,
                cell_range=cell_range,
                rows=self._format_row_range(rows[0], rows[-1]),
            )

        # 第二种：只有一行而且不是键值对，作为普通段落保存。
        if len(rows) == 1:
            return ParsedBlock(
                content=self._build_paragraph_text(
                    reader,
                    rows[0],
                    first_column,
                    last_column,
                ),
                content_type="paragraph",
                section=current_section,
                sheet=reader.worksheet.title,
                cell_range=cell_range,
                rows=str(rows[0]),
            )

        # 第三种：只有闭合边框矩形才按表格处理。
        is_table = self._has_closed_table_border(
            reader,
            rows,
            first_column,
            last_column,
        )
        if is_table:
            header_row_count = self._detect_header_row_count(
                reader,
                rows,
                first_column,
                last_column,
            )
            data_rows = rows[header_row_count:]
            if not data_rows:
                data_rows = rows

            return ParsedBlock(
                content=self._build_table_text(
                    reader,
                    rows,
                    first_column,
                    last_column,
                ),
                content_type="table",
                section=current_section,
                sheet=reader.worksheet.title,
                cell_range=cell_range,
                rows=self._format_row_range(data_rows[0], data_rows[-1]),
            )

        # 多行并不一定是表格。没有闭合边框时，保留为普通段落。
        return ParsedBlock(
            content=self._build_paragraph_region_text(
                reader,
                rows,
                first_column,
                last_column,
            ),
            content_type="paragraph",
            section=current_section,
            sheet=reader.worksheet.title,
            cell_range=cell_range,
            rows=self._format_row_range(rows[0], rows[-1]),
        )

    def _has_closed_table_border(
        self,
        reader: _WorksheetReader,
        rows: list[int],
        first_column: int,
        last_column: int,
    ) -> bool:
        """判断候选区域是否由可见边框围成矩形。 / Checks whether visible borders close the candidate rectangle."""

        first_row = rows[0]
        last_row = rows[-1]
        row_count = last_row - first_row + 1
        column_count = last_column - first_column + 1
        cell_count = row_count * column_count

        if cell_count <= 2:
            return False

        for column_number in range(first_column, last_column + 1):
            top_cell = reader.worksheet.cell(
                row=first_row,
                column=column_number,
            )
            if not self._has_visible_border(top_cell, "top"):
                return False

            bottom_cell = reader.worksheet.cell(
                row=last_row,
                column=column_number,
            )
            if not self._has_visible_border(bottom_cell, "bottom"):
                return False

        for row_number in range(first_row, last_row + 1):
            left_cell = reader.worksheet.cell(
                row=row_number,
                column=first_column,
            )
            if not self._has_visible_border(left_cell, "left"):
                return False

            right_cell = reader.worksheet.cell(
                row=row_number,
                column=last_column,
            )
            if not self._has_visible_border(right_cell, "right"):
                return False

        return True

    @staticmethod
    def _has_visible_border(cell: Cell, side_name: str) -> bool:
        """检查单元格指定方向是否存在可见边框。 / Checks whether one cell side has a visible border."""

        side = getattr(cell.border, side_name)
        if side is None:
            return False

        return side.style is not None

    def _attach_table_context(
        self,
        blocks: list[ParsedBlock],
    ) -> list[ParsedBlock]:
        """把相邻标题、说明和备注加入表格检索正文。 / Adds adjacent titles, descriptions, and notes to table search text."""

        enriched_blocks: list[ParsedBlock] = []

        for block_index in range(len(blocks)):
            block = blocks[block_index]
            if block.content_type != "table":
                enriched_blocks.append(block)
                continue

            context_parts: list[str] = []

            if block_index > 0:
                previous_block = blocks[block_index - 1]
                if self._is_table_context_before(previous_block, block):
                    context_parts.append(
                        f"表格前置说明 / Table context before: {previous_block.content}"
                    )

            context_parts.append(block.content)

            next_index = block_index + 1
            if next_index < len(blocks):
                next_block = blocks[next_index]
                if self._is_table_note_after(block, next_block):
                    context_parts.append(
                        f"表格备注 / Table note: {next_block.content}"
                    )

            enriched_content = "\n".join(context_parts)
            enriched_block = block.model_copy(update={"content": enriched_content})
            enriched_blocks.append(enriched_block)

        return enriched_blocks

    def _is_table_context_before(
        self,
        candidate: ParsedBlock,
        table: ParsedBlock,
    ) -> bool:
        """判断前一个内容块是否是表格的标题或说明。 / Checks whether the previous block describes the table."""

        if candidate.sheet != table.sheet:
            return False

        if candidate.content_type not in ("title", "paragraph"):
            return False

        return self._blocks_are_nearby(candidate, table)

    def _is_table_note_after(
        self,
        table: ParsedBlock,
        candidate: ParsedBlock,
    ) -> bool:
        """判断后一个内容块是否是表格的补充备注。 / Checks whether the following block is a table note."""

        if candidate.sheet != table.sheet:
            return False

        if candidate.content_type != "note":
            return False

        return self._blocks_are_nearby(table, candidate)

    def _blocks_are_nearby(
        self,
        first_block: ParsedBlock,
        second_block: ParsedBlock,
    ) -> bool:
        """判断两个内容块之间是否最多只有一行空白。 / Checks whether blocks have at most one blank row between them."""

        first_last_row = self._last_row_number(first_block.rows)
        second_first_row = self._first_row_number(second_block.rows)
        if first_last_row is None or second_first_row is None:
            return False

        row_distance = second_first_row - first_last_row
        return 0 < row_distance <= 2

    @staticmethod
    def _first_row_number(rows: str | None) -> int | None:
        """读取Block位置中的第一行行号。 / Reads the first row number from block location metadata."""

        if rows is None:
            return None

        first_part = rows.split(":", maxsplit=1)[0]
        return int(first_part)

    @staticmethod
    def _last_row_number(rows: str | None) -> int | None:
        """读取Block位置中的最后一行行号。 / Reads the last row number from block location metadata."""

        if rows is None:
            return None

        row_parts = rows.split(":", maxsplit=1)
        return int(row_parts[-1])

    def _is_key_value_region(
        self,
        reader: _WorksheetReader,
        rows: list[int],
    ) -> bool:
        """判断区域是否由成对的 Label 和 Value 组成。 / Checks whether a region contains label/value pairs."""

        first_row_cells = reader.get_real_nonempty_cells(rows[0])

        # 两行两列更可能是普通小表格，而不是 Key-Value 区域。
        if len(rows) < 3 and len(first_row_cells) == 2:
            return False

        for row_number in rows:
            cells = reader.get_real_nonempty_cells(row_number)
            columns = [column_number for column_number, _value in cells]

            # 每一行至少需要一对，而且单元格数量必须是偶数。
            if len(columns) < 2 or len(columns) % 2 != 0:
                return False

            pairs: list[tuple[int, int]] = []
            pair_start = 0
            while pair_start < len(columns):
                left_column = columns[pair_start]
                right_column = columns[pair_start + 1]

                # Label 和 Value 必须位于相邻两列。
                if right_column != left_column + 1:
                    return False

                pairs.append((left_column, right_column))
                pair_start += 2

            # 一行有多组键值对时，各组之间必须留有空列。
            if len(pairs) > 1:
                has_gap_between_pairs = False
                for pair_index in range(len(pairs) - 1):
                    current_right = pairs[pair_index][1]
                    next_left = pairs[pair_index + 1][0]
                    if next_left > current_right + 1:
                        has_gap_between_pairs = True
                        break

                if not has_gap_between_pairs:
                    return False

        return True

    def _build_key_value_text(
        self,
        reader: _WorksheetReader,
        rows: list[int],
    ) -> str:
        """把键值对区域转换为“名称=值”的文本。 / Converts key-value cells into name=value text."""

        lines: list[str] = []

        for row_number in rows:
            cells = reader.get_real_nonempty_cells(row_number)
            pairs: list[str] = []

            cell_index = 0
            while cell_index < len(cells):
                label = self._to_text(cells[cell_index][1])
                value = self._to_text(cells[cell_index + 1][1])
                pairs.append(f"{label}={value}")
                cell_index += 2

            lines.append("; ".join(pairs))

        return "\n".join(lines)

    def _build_table_text(
        self,
        reader: _WorksheetReader,
        rows: list[int],
        first_column: int,
        last_column: int,
    ) -> str:
        """把表格转换为带完整表头的逐行文本。 / Converts a table into row text with full headers."""

        header_row_count = self._detect_header_row_count(
            reader,
            rows,
            first_column,
            last_column,
        )
        header_rows = rows[:header_row_count]
        data_rows = rows[header_row_count:]

        headers: list[str] = []
        for column_number in range(first_column, last_column + 1):
            header = self._build_header_name(
                reader,
                header_rows,
                column_number,
                first_column,
            )
            headers.append(header)

        if not data_rows:
            return "; ".join(headers)

        lines: list[str] = []

        for row_number in data_rows:
            row_pairs: list[str] = []

            for column_number in range(first_column, last_column + 1):
                value = reader.get_value(row_number, column_number)
                if value in _EMPTY_VALUES:
                    continue

                header_index = column_number - first_column
                header = headers[header_index]
                row_pairs.append(f"{header}={self._to_text(value)}")

            if row_pairs:
                lines.append("; ".join(row_pairs))

        return "\n".join(lines)

    def _detect_header_row_count(
        self,
        reader: _WorksheetReader,
        rows: list[int],
        first_column: int,
        last_column: int,
    ) -> int:
        """判断表格使用一行表头还是两行表头。 / Detects whether a table has one or two header rows."""

        if len(rows) < 3:
            return 1

        first_row = rows[0]
        second_row = rows[1]
        if second_row != first_row + 1:
            return 1

        for merged_range in reader.merged_ranges:
            is_outside_table = (
                merged_range.max_col < first_column
                or merged_range.min_col > last_column
            )
            if is_outside_table or merged_range.min_row != first_row:
                continue

            # 纵向跨到第二行的合并单元格表示两行表头。
            if merged_range.max_row >= second_row:
                return 2

            # 第一行横向合并，且第二行在对应列中有子表头。
            is_horizontal_merge = merged_range.max_col > merged_range.min_col
            if is_horizontal_merge:
                for column_number in range(
                    merged_range.min_col,
                    merged_range.max_col + 1,
                ):
                    second_row_value = reader.get_value(
                        second_row,
                        column_number,
                    )
                    if second_row_value not in _EMPTY_VALUES:
                        return 2

        return 1

    def _build_header_name(
        self,
        reader: _WorksheetReader,
        header_rows: list[int],
        column_number: int,
        first_column: int,
    ) -> str:
        """把父表头和子表头组合成一个完整名称。 / Combines parent and child headers into one name."""

        header_parts: list[str] = []

        for row_number in header_rows:
            value = reader.get_value(row_number, column_number)
            if value in _EMPTY_VALUES:
                continue

            text = self._to_text(value)
            is_duplicate = bool(header_parts) and header_parts[-1] == text
            if not is_duplicate:
                header_parts.append(text)

        if header_parts:
            return " / ".join(header_parts)

        relative_column_number = column_number - first_column + 1
        return f"column_{relative_column_number}"

    def _build_paragraph_text(
        self,
        reader: _WorksheetReader,
        row_number: int,
        first_column: int,
        last_column: int,
    ) -> str:
        """把无法继续分类的一行转换为普通段落。 / Converts an unclassified row into a paragraph."""

        values: list[str] = []

        for column_number in range(first_column, last_column + 1):
            value = reader.get_value(row_number, column_number)
            if value not in _EMPTY_VALUES:
                values.append(self._to_text(value))

        return "; ".join(values)

    def _build_paragraph_region_text(
        self,
        reader: _WorksheetReader,
        rows: list[int],
        first_column: int,
        last_column: int,
    ) -> str:
        """把连续的无边框普通行合并为段落。 / Combines consecutive unbordered rows into a paragraph."""

        lines: list[str] = []

        for row_number in rows:
            line = self._build_paragraph_text(
                reader,
                row_number,
                first_column,
                last_column,
            )
            if line:
                lines.append(line)

        return "\n".join(lines)

    @staticmethod
    def _to_text(value: Any) -> str:
        """把单元格值转换成清理过首尾空白的文本。 / Converts a cell value into trimmed text."""

        return str(value).strip()

    @staticmethod
    def _format_cell_range(
        first_column: int,
        first_row: int,
        last_column: int,
        last_row: int,
    ) -> str:
        """把数字坐标转换成 Excel 的 A1:C3 格式。 / Formats numeric bounds as an Excel A1:C3 range."""

        start = f"{get_column_letter(first_column)}{first_row}"
        end = f"{get_column_letter(last_column)}{last_row}"
        return f"{start}:{end}"

    @staticmethod
    def _format_row_range(first_row: int, last_row: int) -> str:
        """把行号转换成“3”或“3:8”格式。 / Formats rows as either 3 or 3:8."""

        if first_row == last_row:
            return str(first_row)

        return f"{first_row}:{last_row}"
