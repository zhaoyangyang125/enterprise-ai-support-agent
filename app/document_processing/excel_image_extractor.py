"""从 Excel 中提取 openpyxl 能读取的嵌入图片。 / Extracts embedded Excel images readable by openpyxl."""

from pathlib import Path
import logging
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.document_processing.storage import LocalImageAssetStorage


class ExtractedExcelImage:
    """记录一张已保存 Excel 图片的内部来源信息。 / Records internal source information for one stored Excel image."""

    def __init__(
        self,
        image_id: str,
        image_path: Path,
        image_index: int,
        mime_type: str,
        sheet: str,
        cell_range: str | None,
    ) -> None:
        """保存图片资产信息及可靠的 Sheet/锚点定位。 / Stores image asset data and reliable sheet/anchor location.

        参数说明 / Args:
            image_id: 系统生成的稳定图片编号。
            image_path: 图片在服务器内部的保存路径。
            image_index: 图片在整个 Workbook 中的顺序，从 1 开始。
            mime_type: 图片的实际媒体类型。
            sheet: 图片所在的工作表名称。
            cell_range: 能可靠取得时保存图片锚点范围；无法确定时为 None。
        """

        self.image_id = image_id
        self.image_path = image_path
        self.image_index = image_index
        self.mime_type = mime_type
        self.sheet = sheet
        self.cell_range = cell_range


class ExcelImageExtractor:
    """遍历 Excel Sheet，并把嵌入图片交给本地图片存储。 / Walks Excel sheets and sends embedded images to local image storage."""

    def __init__(self, image_storage: LocalImageAssetStorage) -> None:
        """接收 Phase 2 提供的图片资产存储。 / Receives the image asset storage provided by Phase 2.

        参数说明 / Args:
            image_storage: 负责生成 image_id 并保存图片的统一存储对象。
        """

        self._image_storage = image_storage
        self.total_images = 0
        self.failed_images = 0

    def extract(
        self,
        path: Path,
        document_id: str,
        document_version_id: str,
    ) -> list[ExtractedExcelImage]:
        """提取并保存工作簿中的普通嵌入图片。 / Extracts and stores ordinary embedded images from a workbook.

        参数说明 / Args:
            path: 需要读取的 Excel 文件路径。
            document_id: 图片所属文档编号，传给图片存储层。
            document_version_id: 图片所属版本编号，保证不同版本互相隔离。
        """

        # 输入：Excel 路径、文档 ID 和版本 ID。
        # 输出：按 Workbook/Sheet 顺序排列的已保存图片信息。
        # 步骤：打开 Workbook -> 遍历 Sheet -> 读取图片 -> 保存 -> 记录来源。
        workbook = load_workbook(path, data_only=True)
        self.total_images = 0
        self.failed_images = 0
        extracted_images: list[ExtractedExcelImage] = []
        next_image_index = 1

        try:
            for worksheet in workbook.worksheets:
                sheet_images = self._extract_sheet_images(
                    worksheet,
                    document_id,
                    document_version_id,
                    next_image_index,
                )
                extracted_images.extend(sheet_images)
                next_image_index = self.total_images + 1
        finally:
            workbook.close()

        return extracted_images

    def _extract_sheet_images(
        self,
        worksheet: Worksheet,
        document_id: str,
        document_version_id: str,
        first_image_index: int,
    ) -> list[ExtractedExcelImage]:
        """按 Sheet 内部顺序提取图片并保存来源定位。 / Extracts sheet images in order and stores source locations.

        参数说明 / Args:
            worksheet: 当前正在处理的 openpyxl 工作表对象。
            document_id: 图片所属文档编号。
            document_version_id: 图片所属版本编号。
            first_image_index: 当前 Sheet 第一张图片应使用的全 Workbook 序号。
        """

        extracted_images: list[ExtractedExcelImage] = []

        # openpyxl 3.1 没有公开的 Worksheet 图片迭代接口。
        # `_images` 是其读取普通嵌入图片时实际维护的集合；SmartArt、Shape、Chart 不在这里。
        embedded_images = getattr(worksheet, "_images", [])
        self.total_images += len(embedded_images)

        for offset in range(len(embedded_images)):
            embedded_image = embedded_images[offset]
            image_index = first_image_index + offset
            try:
                image_content = embedded_image._data()
                mime_type = self._get_output_mime_type(embedded_image)
                stored_asset = self._image_storage.store(
                    content=image_content,
                    document_id=document_id,
                    document_version_id=document_version_id,
                    image_index=image_index,
                    mime_type=mime_type,
                )
            except Exception:
                self.failed_images += 1
                logging.getLogger(__name__).warning(
                    "image_extraction_failed version=%s image=%s",
                    document_version_id, image_index,
                )
                continue
            try:
                cell_range = self._get_anchor_cell_range(embedded_image)
            except Exception:
                # 锚点不可靠时保留图片，但不伪造位置。
                cell_range = None
            extracted_image = ExtractedExcelImage(
                image_id=stored_asset.image_id,
                image_path=stored_asset.path,
                image_index=stored_asset.image_index,
                mime_type=stored_asset.mime_type,
                sheet=worksheet.title,
                cell_range=cell_range,
            )
            extracted_images.append(extracted_image)

        return extracted_images

    @staticmethod
    def _get_output_mime_type(embedded_image: Any) -> str:
        """返回 `_data()` 实际输出图片格式对应的 MIME。 / Returns the MIME type produced by openpyxl `_data()`.

        参数说明 / Args:
            embedded_image: openpyxl 读取到的内部图片对象，用于检查实际输出格式。
        """

        image_format = getattr(embedded_image, "format", "")
        normalized_format = str(image_format).casefold()

        if normalized_format == "jpeg" or normalized_format == "jpg":
            return "image/jpeg"
        if normalized_format == "gif":
            return "image/gif"
        if normalized_format == "png":
            return "image/png"

        # openpyxl 会把 BMP 等其他可读取格式转换成 PNG 后再返回 bytes。
        return "image/png"

    @staticmethod
    def _get_anchor_cell_range(embedded_image: Any) -> str | None:
        """只在 openpyxl 提供可靠单元格锚点时返回定位。 / Returns a location only when openpyxl provides reliable cell anchors.

        参数说明 / Args:
            embedded_image: Excel 嵌入图片对象，从中读取起点和终点锚点。
        """

        anchor = getattr(embedded_image, "anchor", None)
        if anchor is None:
            return None
        if isinstance(anchor, str):
            return anchor

        start_marker = getattr(anchor, "_from", None)
        if start_marker is None:
            return None

        start_cell = ExcelImageExtractor._marker_to_cell(start_marker)
        end_marker = getattr(anchor, "to", None)
        if end_marker is None:
            return start_cell

        end_cell = ExcelImageExtractor._marker_to_cell(end_marker)
        if end_cell == start_cell:
            return start_cell
        return f"{start_cell}:{end_cell}"

    @staticmethod
    def _marker_to_cell(marker: Any) -> str:
        """把 openpyxl 的零起点锚点转换成 Excel 单元格地址。 / Converts a zero-based openpyxl marker to an Excel cell address.

        参数说明 / Args:
            marker: openpyxl 的图片锚点标记，列号和行号都从 0 开始。
        """

        column_number = int(marker.col) + 1
        row_number = int(marker.row) + 1
        column_name = get_column_letter(column_number)
        return f"{column_name}{row_number}"
