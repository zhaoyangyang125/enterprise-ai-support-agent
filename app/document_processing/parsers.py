from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook
from pypdf import PdfReader

from app.schemas.document import ParsedBlock


class DocumentParser(Protocol):
    """定义原始文档转换为带定位语义块的接口。 / Defines the interface for converting an original document into located semantic blocks."""

    def parse(self, path: Path) -> list[ParsedBlock]:
        """解析指定本地文档。 / Parses the specified local document."""

        ...


class PdfDocumentParser:
    """按页和自然段解析带文本层的 PDF。 / Parses text-layer PDFs by page and paragraph."""

    def __init__(self, maximum_chunk_characters: int = 1200) -> None:
        """设置语义段落合并后的最大字符数。 / Sets the maximum size after combining semantic paragraphs."""

        self._maximum_chunk_characters = maximum_chunk_characters

    def parse(self, path: Path) -> list[ParsedBlock]:
        """保留页码并将连续自然段组合成 Chunk。 / Preserves page numbers and groups consecutive paragraphs into chunks."""

        blocks: list[ParsedBlock] = []
        for page_number, page in enumerate(PdfReader(path).pages, start=1):
            text = page.extract_text() or ""
            paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
            current: list[str] = []
            current_length = 0
            for paragraph in paragraphs:
                if current and current_length + len(paragraph) > self._maximum_chunk_characters:
                    blocks.append(
                        ParsedBlock(content="\n\n".join(current), page=page_number)
                    )
                    current = []
                    current_length = 0
                current.append(paragraph)
                current_length += len(paragraph)
            if current:
                blocks.append(ParsedBlock(content="\n\n".join(current), page=page_number))
        return blocks


class ExcelDocumentParser:
    """按 Sheet 结构和行关系解析 Excel，而不是固定字符切割。 / Parses Excel by sheet structure and row relationships instead of fixed character slicing."""

    def parse(self, path: Path) -> list[ParsedBlock]:
        """展开合并单元格，并把每个数据行转换为 Header=Value 语义记录。 / Expands merged cells and converts each data row into a Header=Value semantic record."""

        workbook = load_workbook(path, data_only=True)
        blocks: list[ParsedBlock] = []
        for worksheet in workbook.worksheets:
            self._expand_merged_cells(worksheet)
            populated_rows = [
                (row_index, [cell.value for cell in row])
                for row_index, row in enumerate(worksheet.iter_rows(), start=1)
                if any(cell.value not in (None, "") for cell in row)
            ]
            if not populated_rows:
                continue
            header_row_index, header_values = populated_rows[0]
            headers = [
                str(value).strip() if value not in (None, "") else f"column_{index}"
                for index, value in enumerate(header_values, start=1)
            ]
            data_rows = populated_rows[1:] or [(header_row_index, header_values)]
            for row_index, values in data_rows:
                pairs = [
                    f"{headers[index]}={value}"
                    for index, value in enumerate(values)
                    if value not in (None, "")
                ]
                if pairs:
                    blocks.append(
                        ParsedBlock(
                            content="; ".join(pairs),
                            sheet=worksheet.title,
                            rows=str(row_index),
                        )
                    )
        workbook.close()
        return blocks

    @staticmethod
    def _expand_merged_cells(worksheet: object) -> None:
        """把合并区域左上角的值复制到区域内所有单元格。 / Copies each merged region's top-left value to every cell in that region."""

        ranges = list(worksheet.merged_cells.ranges)
        for merged_range in ranges:
            value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
            worksheet.unmerge_cells(str(merged_range))
            for row in worksheet.iter_rows(
                min_row=merged_range.min_row,
                max_row=merged_range.max_row,
                min_col=merged_range.min_col,
                max_col=merged_range.max_col,
            ):
                for cell in row:
                    cell.value = value


class DocumentParserRegistry:
    """根据扩展名选择受支持的本地文档解析器。 / Selects a supported local document parser by file extension."""

    def __init__(self) -> None:
        """注册 v1 支持的 PDF 和 Excel 解析器。 / Registers the PDF and Excel parsers supported in v1."""

        excel_parser = ExcelDocumentParser()
        self._parsers: dict[str, DocumentParser] = {
            ".pdf": PdfDocumentParser(),
            ".xlsx": excel_parser,
            ".xlsm": excel_parser,
        }

    def get(self, path: Path) -> DocumentParser:
        """返回对应解析器，不支持的格式会明确失败。 / Returns the matching parser and explicitly fails for unsupported formats."""

        parser = self._parsers.get(path.suffix.casefold())
        if parser is None:
            raise ValueError(f"Unsupported document format: {path.suffix}")
        return parser
